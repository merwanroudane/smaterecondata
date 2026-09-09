"""Excel workbook export (spec 15, 0L).

XLSX is a flagship output, not an afterthought. A workbook produced here is
meant to be handed straight to a co-author: every column can be traced back to
a provider series, every transformation is listed, and the request that
produced it is recorded.

Sheets: Data, Metadata, Descriptive_Stats, Missing_Data, Sources,
Transformations, Query, README.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..analytics import missingness as missingness_mod
from ..analytics.descriptive import DescriptiveStats, describe
from ..core.models import DataQualityReport, OutputShape
from ..datasets.builder import Dataset

# Sunrise Research palette (spec 0G), muted for print.
_HEADER_FILL = PatternFill("solid", fgColor="F26B4F")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_TITLE_FONT = Font(bold=True, size=14, color="25313C")
_SUBTLE_FONT = Font(color="66727D", size=10)
_WARN_FILL = PatternFill("solid", fgColor="FDF0D5")
_ERROR_FILL = PatternFill("solid", fgColor="FBE3E3")
_MISSING_FILL = PatternFill("solid", fgColor="F4F4F4")

_THIN = Side(style="thin", color="E9E3DC")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# Excel forbids these in a sheet name, and caps names at 31 characters.
_ILLEGAL_SHEET = set(r"[]:*?/\\")

MAX_ROWS_PER_SHEET = 1_000_000  # Excel's own limit is 1,048,576


@dataclass
class ExcelExportResult:
    path: Path
    sheets: list[str]
    rows: int
    truncated: bool = False


def safe_sheet_name(name: str) -> str:
    cleaned = "".join("_" if ch in _ILLEGAL_SHEET else ch for ch in name)
    return cleaned[:31] or "Sheet"


def _number_format(values: Sequence[Any]) -> str:
    """Pick a readable format from the magnitude of the data."""
    numeric = [abs(float(v)) for v in values
               if isinstance(v, (int, float)) and v is not None]
    if not numeric:
        return "General"
    largest = max(numeric)
    if largest >= 1_000_000:
        return "#,##0"
    if largest >= 1_000:
        return "#,##0.0"
    if largest >= 1:
        return "#,##0.00"
    return "0.0000"


def _write_header(ws: Worksheet, headers: Sequence[str], row: int = 1) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = _BORDER


def _autosize(ws: Worksheet, headers: Sequence[str],
              rows: Sequence[dict], *, minimum: int = 9,
              maximum: int = 52) -> None:
    for col, key in enumerate(headers, start=1):
        widest = len(str(key))
        for row in rows[:400]:  # sampling keeps very large exports fast
            value = row.get(key)
            if value is not None:
                widest = max(widest, len(str(value)))
        ws.column_dimensions[get_column_letter(col)].width = min(
            max(widest + 2, minimum), maximum
        )


def _write_table(ws: Worksheet, headers: Sequence[str],
                 rows: Sequence[dict], *, freeze: str = "A2",
                 autofilter: bool = True,
                 highlight_missing: bool = False) -> None:
    """Write a plain rectangular table -- never a merged cell inside it."""
    _write_header(ws, headers)

    column_values: dict[str, list[Any]] = {h: [] for h in headers}
    for r_index, row in enumerate(rows, start=2):
        for c_index, key in enumerate(headers, start=1):
            value = row.get(key)
            if isinstance(value, (list, tuple)):
                value = ", ".join(str(v) for v in value)
            elif isinstance(value, _dt.datetime):
                value = value.replace(tzinfo=None)
            cell = ws.cell(row=r_index, column=c_index, value=value)
            cell.border = _BORDER
            if value is None and highlight_missing:
                cell.fill = _MISSING_FILL
            column_values[key].append(value)

    # One number format per column, from that column's own magnitudes.
    for c_index, key in enumerate(headers, start=1):
        fmt = _number_format(column_values[key])
        if fmt == "General":
            continue
        letter = get_column_letter(c_index)
        for r_index in range(2, len(rows) + 2):
            ws[f"{letter}{r_index}"].number_format = fmt

    if freeze:
        ws.freeze_panes = freeze
    if autofilter and rows:
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        )
    _autosize(ws, headers, rows)


def _write_kv(ws: Worksheet, title: str,
              pairs: Sequence[tuple[str, Any]], start: int = 1) -> int:
    """Key/value block. Returns the next free row."""
    ws.cell(row=start, column=1, value=title).font = _TITLE_FONT
    row = start + 1
    for key, value in pairs:
        k = ws.cell(row=row, column=1, value=key)
        k.font = Font(bold=True, size=10)
        if isinstance(value, (list, tuple)):
            value = ", ".join(str(v) for v in value)
        elif isinstance(value, _dt.datetime):
            value = value.replace(tzinfo=None)
        ws.cell(row=row, column=2, value=value).alignment = Alignment(
            wrap_text=True, vertical="top")
        row += 1
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 92
    return row + 1


def export_workbook(
    dataset: Dataset,
    path: str | Path,
    *,
    quality: DataQualityReport | None = None,
    stats: Sequence[DescriptiveStats] | None = None,
    platform_version: str = "SmatEconData 1.0",
) -> ExcelExportResult:
    """Write the full SmatEconData workbook."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    lineage = dataset.lineage or dataset.build_lineage()
    miss = missingness_mod.analyse(dataset)

    if stats is None:
        columns = dataset.wide_columns_map()
        _, wide_rows = dataset.to_wide()
        periods = [str(r["period"]) for r in wide_rows]
        units = {v.alias: v.metadata.unit for v in dataset.variables}
        stats = [
            describe(values, variable=alias, periods=periods,
                     unit=units.get(alias))
            for alias, values in columns.items()
        ]

    truncated = False

    # -- 1. Data ---------------------------------------------------------
    ws = wb.create_sheet(safe_sheet_name("Data"))
    if dataset.spec.output_shape is OutputShape.LONG:
        data_rows = dataset.to_long()
        headers = ["geography", "iso3", "period", "indicator", "value", "unit",
                   "frequency", "provider", "series_id", "status"]
    else:
        headers, data_rows = dataset.to_wide()

    if len(data_rows) > MAX_ROWS_PER_SHEET:
        data_rows = data_rows[:MAX_ROWS_PER_SHEET]
        truncated = True
    # Freeze the identifier columns as well as the header row.
    _write_table(ws, headers, data_rows, freeze="D2", highlight_missing=True)

    # -- 2. Metadata -----------------------------------------------------
    ws = wb.create_sheet(safe_sheet_name("Metadata"))
    meta_headers = ["variable_name", "provider", "series_id", "series_title",
                    "definition", "unit", "frequency", "coverage_start",
                    "coverage_end", "retrieved_at", "source_reference", "notes"]
    meta_rows = [{
        "variable_name": line.column,
        "provider": line.provider,
        "series_id": line.series_id,
        "series_title": line.series_title,
        "definition": line.definition,
        "unit": line.unit,
        "frequency": line.frequency.value,
        "coverage_start": line.coverage_start,
        "coverage_end": line.coverage_end,
        "retrieved_at": line.retrieved_at,
        "source_reference": line.source_reference,
        "notes": "; ".join(line.warnings) if line.warnings else None,
    } for line in lineage]
    _write_table(ws, meta_headers, meta_rows)

    # -- 3. Descriptive_Stats --------------------------------------------
    ws = wb.create_sheet(safe_sheet_name("Descriptive_Stats"))
    stat_rows = [s.as_row() for s in stats]
    stat_headers = list(stat_rows[0]) if stat_rows else ["variable"]
    _write_table(ws, stat_headers, stat_rows)

    # -- 4. Missing_Data --------------------------------------------------
    ws = wb.create_sheet(safe_sheet_name("Missing_Data"))
    miss_rows = miss.as_rows()
    miss_headers = [
        "scope", "name", "observations", "present", "missing", "missing_pct",
        "coverage_pct", "first_valid", "last_valid", "longest_gap",
        "longest_gap_span", "longest_gap_geography",
    ]
    _write_table(ws, miss_headers, miss_rows)

    # -- 5. Sources -------------------------------------------------------
    ws = wb.create_sheet(safe_sheet_name("Sources"))
    source_headers = ["variable_name", "provider", "series_id", "source_reference",
                      "retrieved_at", "citation"]
    source_rows = [{
        "variable_name": line.column,
        "provider": line.provider,
        "series_id": line.series_id,
        "source_reference": line.source_reference,
        "retrieved_at": line.retrieved_at,
        "citation": line.citation,
    } for line in lineage]
    _write_table(ws, source_headers, source_rows)
    # Hyperlink the reference column only when it is a real http(s) URL.
    ref_col = source_headers.index("source_reference") + 1
    for r_index, row in enumerate(source_rows, start=2):
        ref = row.get("source_reference")
        if isinstance(ref, str) and ref.startswith(("http://", "https://")):
            cell = ws.cell(row=r_index, column=ref_col)
            cell.hyperlink = ref
            cell.font = Font(color="0563C1", underline="single")

    # -- 6. Transformations ------------------------------------------------
    ws = wb.create_sheet(safe_sheet_name("Transformations"))
    trans_headers = ["step", "operation", "parameters", "input_columns",
                     "output_columns", "timestamp", "lossy", "note"]
    trans_rows = [{
        "step": i,
        "operation": t.operation,
        "parameters": "; ".join(f"{k}={v}" for k, v in t.parameters.items()),
        "input_columns": ", ".join(t.input_columns),
        "output_columns": ", ".join(t.output_columns),
        "timestamp": t.timestamp,
        "lossy": "yes" if t.lossy else "no",
        "note": t.note,
    } for i, t in enumerate(dataset.transformations, start=1)]
    _write_table(ws, trans_headers, trans_rows)

    # -- 7. Query ----------------------------------------------------------
    ws = wb.create_sheet(safe_sheet_name("Query"))
    spec = dataset.spec
    row = _write_kv(ws, "Original request", [
        ("Query", spec.original_query or "(built programmatically)"),
        ("Query language", spec.query_language.value if spec.query_language else "n/a"),
    ])
    row = _write_kv(ws, "Resolved DatasetSpec", [
        ("Geographies", spec.geographies),
        ("Indicators", [i.concept for i in spec.indicators]),
        ("Period", spec.period_label()),
        ("Frequency", spec.frequency.value),
        ("Output shape", spec.output_shape.value),
        ("Preferred sources", spec.preferred_sources or ["automatic"]),
        ("Export formats", spec.export_formats),
    ], start=row)
    if quality is not None:
        row = _write_kv(ws, "Data quality", [
            ("Status", quality.status),
            ("Rows", quality.rows),
            ("Countries", quality.countries),
            ("Indicators", quality.indicators),
            ("Coverage %", quality.coverage_pct),
            ("Duplicate keys", quality.duplicate_keys),
            ("Missing cells", quality.missing_cells),
            ("Frequency conflicts", quality.frequency_conflicts),
            ("Unit conflicts", quality.unit_conflicts),
            ("Warnings", quality.warnings),
            ("Errors", quality.errors),
        ], start=row)

        ws.cell(row=row, column=1, value="Validation issues").font = _TITLE_FONT
        row += 1
        _write_header(ws, ["severity", "code", "variable", "geography", "message"],
                      row=row)
        row += 1
        for issue in quality.issues:
            ws.cell(row=row, column=1, value=issue.severity.value)
            ws.cell(row=row, column=2, value=issue.code)
            ws.cell(row=row, column=3, value=issue.variable)
            ws.cell(row=row, column=4, value=issue.geography)
            ws.cell(row=row, column=5, value=issue.message).alignment = Alignment(
                wrap_text=True, vertical="top")
            if issue.severity.value == "error":
                ws.cell(row=row, column=1).fill = _ERROR_FILL
            elif issue.severity.value == "warning":
                ws.cell(row=row, column=1).fill = _WARN_FILL
            row += 1
        ws.column_dimensions["C"].width = 24
        ws.column_dimensions["D"].width = 14
        ws.column_dimensions["E"].width = 90

    # -- 8. README ---------------------------------------------------------
    ws = wb.create_sheet(safe_sheet_name("README"))
    _write_kv(ws, "SmatEconData export", [
        ("Dataset", dataset.name),
        ("Built", dataset.retrieved_at),
        ("Rows", len(data_rows)),
        ("Variables", len(dataset.variables)),
        ("Countries", len(dataset.geographies)),
        ("Period", spec.period_label()),
        ("Platform", platform_version),
        ("Author", "Dr Merwan Roudane"),
        ("Repository", "https://github.com/merwanroudane/smaterecondata"),
        ("", ""),
        ("Sheet: Data", "The dataset itself. Blank cells are genuine gaps; "
                        "nothing was interpolated or imputed."),
        ("Sheet: Metadata", "One row per variable: provider, series id, "
                            "official title, definition, unit, coverage."),
        ("Sheet: Descriptive_Stats", "Descriptive statistics only. Missing "
                                     "values are excluded, never filled."),
        ("Sheet: Missing_Data", "Coverage and gaps by variable and by country."),
        ("Sheet: Sources", "Provenance and a ready-to-paste citation per column."),
        ("Sheet: Transformations", "Every operation applied, in order."),
        ("Sheet: Query", "The original request, the resolved specification, "
                         "and the data-quality report."),
    ])

    wb.save(path)
    return ExcelExportResult(
        path=path,
        sheets=wb.sheetnames,
        rows=len(data_rows),
        truncated=truncated,
    )
