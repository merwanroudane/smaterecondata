"""Dataset assembly, validation, missingness and export tests.

Covers spec sections 7 (builder), 8 (validation), 9 (missing data),
15 (Excel export) and 16 (HTML report).

Import-light on purpose: no provider, no network, no ``curl_cffi``.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from backend.smatecondata.analytics import missingness as missingness_mod
from backend.smatecondata.analytics.descriptive import describe
from backend.smatecondata.core.models import (
    DatasetSpec,
    Frequency,
    IndicatorMetadata,
    Observation,
    OutputShape,
    Severity,
)
from backend.smatecondata.datasets.builder import (
    DatasetBuilder,
    dedupe_aliases,
    smart_alias,
)
from backend.smatecondata.datasets.validation import (
    _modified_z_scores,
    infer_frequency,
    period_sort_key,
    unit_family,
    validate_dataset,
)
from backend.smatecondata.exports.excel import export_workbook, safe_sheet_name
from backend.smatecondata.exports.html_report import export_report

GDP = IndicatorMetadata(
    provider="world_bank", series_id="NY.GDP.PCAP.CD",
    title="GDP per capita (current US$)", unit="current US$",
    frequency=Frequency.ANNUAL, source_name="World Bank",
    source_reference="https://data.worldbank.org/indicator/NY.GDP.PCAP.CD")

CPI = IndicatorMetadata(
    provider="imf", series_id="PCPIPCH",
    title="Inflation, average consumer prices", unit="Annual %",
    frequency=Frequency.ANNUAL, source_name="IMF")

NAMES = {"DZA": "Algeria", "MAR": "Morocco"}


def obs(metadata, iso3, period, value, status=None):
    return Observation(
        provider=metadata.provider, series_id=metadata.series_id,
        geography=NAMES.get(iso3, iso3), iso3=iso3, period=str(period),
        value=value, unit=metadata.unit, frequency=Frequency.ANNUAL,
        status=status)


@pytest.fixture
def spec():
    return DatasetSpec(geographies=["DZA", "MAR"], start_year=2000,
                       end_year=2005, frequency=Frequency.ANNUAL)


@pytest.fixture
def dataset(spec):
    """Two variables with deliberately planted data problems."""
    gdp = ([obs(GDP, "DZA", y, 1800.0 + (y - 2000) * 120) for y in range(2000, 2006)]
           + [obs(GDP, "MAR", y, 1500.0 + (y - 2000) * 100) for y in range(2000, 2006)])
    # DZA: interior gap at 2002-2003.  MAR: spike at 2003.
    cpi = [obs(CPI, "DZA", y, 2.5 + (y - 2000) * 0.3) for y in (2000, 2001, 2004, 2005)]
    cpi += [obs(CPI, "MAR", y, 1.9 if y != 2003 else 45.0) for y in range(2000, 2006)]
    return (DatasetBuilder(spec, name="test_dataset")
            .add_series(GDP, gdp, "gdp_per_capita")
            .add_series(CPI, cpi, "inflation")
            .build())


# --------------------------------------------------------------------------
# Column naming (spec 43)
# --------------------------------------------------------------------------


def test_smart_alias_is_readable_not_a_provider_code():
    assert smart_alias(GDP, "gdp_per_capita") == "gdp_per_capita_usd"
    assert smart_alias(CPI, "inflation") == "inflation_pct"


def test_aliases_are_deduped_without_dropping_any():
    assert dedupe_aliases(["gdp", "gdp", "cpi"]) == ["gdp", "gdp_2", "cpi"]


def test_sheet_names_are_valid():
    assert safe_sheet_name("Data/2024:v1") == "Data_2024_v1"
    assert len(safe_sheet_name("x" * 60)) == 31


# --------------------------------------------------------------------------
# Shapes (spec 7.1)
# --------------------------------------------------------------------------


def test_wide_shape_has_one_row_per_geography_period(dataset):
    columns, rows = dataset.to_wide()
    assert columns == ["geography", "iso3", "period",
                       "gdp_per_capita_usd", "inflation_pct"]
    assert len(rows) == 12  # 2 countries x 6 years
    assert len({(r["iso3"], r["period"]) for r in rows}) == 12


def test_long_and_wide_are_switchable_without_refetch(dataset):
    assert len(dataset.to_long()) == len(dataset.observations)
    assert len(dataset.to_wide()[1]) == 12
    # Both shapes come from the same in-memory observations.
    assert dataset.to_long()[0]["value"] is not None


def test_gaps_stay_none_and_are_never_imputed(dataset):
    _, rows = dataset.to_wide()
    dza = {r["period"]: r["inflation_pct"] for r in rows if r["iso3"] == "DZA"}
    assert dza["2002"] is None
    assert dza["2003"] is None
    assert dza["2001"] == pytest.approx(2.8)


def test_output_shape_drives_rows(spec, dataset):
    dataset.spec.output_shape = OutputShape.LONG
    assert dataset.rows()[0]["indicator"] in dataset.aliases
    dataset.spec.output_shape = OutputShape.WIDE
    assert "gdp_per_capita_usd" in dataset.rows()[0]


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------


@pytest.mark.parametrize("period,frequency", [
    ("2020", Frequency.ANNUAL),
    ("2020Q3", Frequency.QUARTERLY),
    ("2020-07", Frequency.MONTHLY),
    ("2020-07-15", Frequency.DAILY),
    ("garbage", Frequency.UNKNOWN),
])
def test_frequency_inference(period, frequency):
    assert infer_frequency(period) is frequency


def test_periods_sort_chronologically():
    given = ["2001", "2000", "2010", "1999"]
    assert sorted(given, key=period_sort_key) == ["1999", "2000", "2001", "2010"]


@pytest.mark.parametrize("unit,family", [
    ("current US$", "current_usd"),
    ("constant 2015 US$", "constant_usd"),
    ("Annual %", "percent"),
    ("Index, 2010 = 100", "index"),
    (None, None),
])
def test_unit_families(unit, family):
    assert unit_family(unit) == family


# --------------------------------------------------------------------------
# Validation (spec 8)
# --------------------------------------------------------------------------


def test_gap_is_reported(dataset):
    codes = {i.code for i in validate_dataset(dataset).issues}
    assert "time_gap" in codes


def test_duplicate_key_is_an_error(spec):
    rows = [obs(GDP, "DZA", 2000, 1800.0), obs(GDP, "DZA", 2000, 1801.0)]
    report = validate_dataset(
        DatasetBuilder(spec).add_series(GDP, rows, "gdp_per_capita").build())
    dup = [i for i in report.issues if i.code == "duplicate_key"]
    assert dup and dup[0].severity is Severity.ERROR
    assert report.status == "unusable"


def test_mixed_frequency_is_an_error(spec):
    rows = [obs(GDP, "DZA", "2000", 1.0), obs(GDP, "DZA", "2000Q1", 2.0)]
    codes = {i.code for i in validate_dataset(
        DatasetBuilder(spec).add_series(GDP, rows, "gdp").build()).issues}
    assert "mixed_frequency" in codes


def test_unit_conflict_within_one_variable(spec):
    a = obs(GDP, "DZA", 2000, 1.0)
    b = obs(GDP, "DZA", 2001, 2.0)
    b.unit = "constant 2015 US$"
    codes = {i.code for i in validate_dataset(
        DatasetBuilder(spec).add_series(GDP, [a, b], "gdp").build()).issues}
    assert "unit_conflict" in codes


def test_all_values_missing_is_flagged(spec):
    rows = [obs(GDP, "DZA", y, None) for y in range(2000, 2006)]
    codes = {i.code for i in validate_dataset(
        DatasetBuilder(spec).add_series(GDP, rows, "gdp").build()).issues}
    assert "all_values_missing" in codes


def test_coverage_ends_early_is_reported_per_country(spec):
    """One country stopping early must not be hidden by the others."""
    rows = [obs(GDP, "DZA", y, 100.0 + y) for y in range(2000, 2006)]
    rows += [obs(GDP, "MAR", y, 200.0 + y) for y in range(2000, 2004)]
    report = validate_dataset(
        DatasetBuilder(spec).add_series(GDP, rows, "gdp").build())
    early = [i for i in report.issues if i.code == "coverage_ends_early"]
    assert len(early) == 1
    assert early[0].geography == "MAR"
    assert early[0].details["last_year"] == 2003


def test_extreme_jump_is_info_only_and_changes_nothing(dataset):
    report = validate_dataset(dataset)
    jumps = [i for i in report.issues if i.code == "extreme_jump"]
    assert jumps, "the planted 45% spike should be flagged"
    assert all(i.severity is Severity.INFO for i in jumps)
    # The value itself is untouched.
    _, rows = dataset.to_wide()
    spike = [r for r in rows if r["iso3"] == "MAR" and r["period"] == "2003"]
    assert spike[0]["inflation_pct"] == 45.0


def test_modified_z_score_catches_a_spike_a_stdev_rule_would_miss():
    scores = _modified_z_scores([1.9, 1.9, 1.9, 45.0, 1.9, 1.9])
    assert scores is not None
    assert max(abs(z) for z in scores) > 3.5


def test_modified_z_score_does_not_flag_a_linear_trend():
    scores = _modified_z_scores([1800, 1920, 2040, 2160, 2280, 2400])
    assert scores is not None
    assert max(abs(z) for z in scores) < 3.5


def test_modified_z_score_returns_none_for_a_constant_series():
    assert _modified_z_scores([2.0] * 6) is None


def test_clean_dataset_reports_clean(spec):
    rows = [obs(GDP, "DZA", y, 1800.0 + (y - 2000) * 120) for y in range(2000, 2006)]
    report = validate_dataset(
        DatasetBuilder(spec).add_series(GDP, rows, "gdp_per_capita").build())
    assert report.errors == 0
    assert report.coverage_pct == 100.0


# --------------------------------------------------------------------------
# Missingness (spec 9)
# --------------------------------------------------------------------------


def test_longest_gap_is_measured_per_country(dataset):
    report = missingness_mod.analyse(dataset)
    inflation = next(v for v in report.by_variable if v.variable == "inflation_pct")
    assert inflation.missing == 2
    assert inflation.longest_gap == 2
    assert inflation.longest_gap_span == ("2002", "2003")
    assert inflation.longest_gap_geography == "DZA"


def test_missingness_by_country(dataset):
    report = missingness_mod.analyse(dataset)
    dza = next(g for g in report.by_geography if g.geography == "DZA")
    assert dza.missing == 2
    assert dza.coverage_pct == pytest.approx(83.33, abs=0.01)


def test_coverage_matrix_marks_the_gap(dataset):
    report = missingness_mod.analyse(dataset)
    row = report.coverage_matrix["inflation_pct"]
    assert row["2000"] == 100.0
    assert row["2002"] == 50.0  # MAR present, DZA missing


def test_fill_is_marked_and_never_silent():
    filled, flags = missingness_mod.fill_marked([1.0, None, None, 4.0])
    assert filled == [1.0, 1.0, 1.0, 4.0]
    assert flags == [False, True, True, False]


def test_fill_rejects_an_unknown_method():
    with pytest.raises(ValueError):
        missingness_mod.fill_marked([1.0, None], method="magic")


# --------------------------------------------------------------------------
# Provenance (spec 14, 44)
# --------------------------------------------------------------------------


def test_every_column_carries_lineage(dataset):
    lineage = dataset.build_lineage()
    assert {line.column for line in lineage} == set(dataset.aliases)
    for line in lineage:
        assert line.provider and line.series_id and line.series_title
        assert line.citation and "Retrieved" in line.citation


def test_build_is_recorded_in_the_transformation_log(dataset):
    assert dataset.transformations
    assert dataset.transformations[0].operation == "build_dataset"
    assert set(dataset.transformations[0].output_columns) == set(dataset.aliases)


# --------------------------------------------------------------------------
# Excel export (spec 15, 39)
# --------------------------------------------------------------------------


def test_workbook_has_every_required_sheet(dataset, tmp_path):
    result = export_workbook(dataset, tmp_path / "out.xlsx",
                             quality=validate_dataset(dataset))
    assert result.sheets == ["Data", "Metadata", "Descriptive_Stats",
                             "Missing_Data", "Sources", "Transformations",
                             "Query", "README"]


def test_workbook_opens_and_row_counts_match(dataset, tmp_path):
    path = tmp_path / "out.xlsx"
    export_workbook(dataset, path, quality=validate_dataset(dataset))

    wb = load_workbook(path)
    ws = wb["Data"]
    _, rows = dataset.to_wide()
    assert ws.max_row == len(rows) + 1          # + header
    assert ws.freeze_panes == "D2"
    assert ws.auto_filter.ref is not None

    # Metadata has exactly one row per variable.
    assert wb["Metadata"].max_row == len(dataset.variables) + 1


def test_workbook_keeps_numbers_numeric(dataset, tmp_path):
    path = tmp_path / "out.xlsx"
    export_workbook(dataset, path)
    ws = load_workbook(path)["Data"]
    assert isinstance(ws["D2"].value, (int, float))
    assert ws["D2"].number_format != "General"


def test_workbook_hyperlinks_only_real_urls(dataset, tmp_path):
    path = tmp_path / "out.xlsx"
    export_workbook(dataset, path)
    ws = load_workbook(path)["Sources"]
    links = [ws.cell(row=r, column=4).hyperlink
             for r in range(2, ws.max_row + 1)]
    targets = [l.target for l in links if l is not None]
    assert all(t.startswith("http") for t in targets)
    assert any("worldbank.org" in t for t in targets)


def test_workbook_has_no_merged_cells_in_the_data_table(dataset, tmp_path):
    path = tmp_path / "out.xlsx"
    export_workbook(dataset, path)
    assert not load_workbook(path)["Data"].merged_cells.ranges


# --------------------------------------------------------------------------
# HTML report (spec 16)
# --------------------------------------------------------------------------


def test_html_report_is_standalone_and_complete(dataset, tmp_path):
    path = tmp_path / "report.html"
    export_report(dataset, path, quality=validate_dataset(dataset))
    markup = path.read_text(encoding="utf-8")

    for section in ("Selected indicators", "Charts", "Missing data",
                    "Descriptive statistics", "Sources and provenance",
                    "Transformation history", "Reproducibility"):
        assert section in markup, f"missing section: {section}"

    assert "Dr Merwan Roudane" in markup
    assert "github.com/merwanroudane/smaterecondata" in markup
    # Charts are inline SVG: no external script or stylesheet is fetched.
    assert "<svg" in markup
    assert "<script" not in markup
    assert "cdn." not in markup


def test_html_report_escapes_content(spec, tmp_path):
    nasty = IndicatorMetadata(
        provider="test", series_id="X<1>",
        title="<script>alert('xss')</script>", unit="%",
        frequency=Frequency.ANNUAL)
    rows = [obs(nasty, "DZA", y, float(y)) for y in range(2000, 2006)]
    dataset = DatasetBuilder(spec).add_series(nasty, rows, "gdp").build()
    path = tmp_path / "report.html"
    export_report(dataset, path)
    markup = path.read_text(encoding="utf-8")
    assert "<script>alert" not in markup
    assert "&lt;script&gt;" in markup


def test_html_report_states_that_nothing_was_imputed(dataset, tmp_path):
    path = tmp_path / "report.html"
    export_report(dataset, path, quality=validate_dataset(dataset))
    assert "never imputed" in path.read_text(encoding="utf-8")


# --------------------------------------------------------------------------
# Descriptive statistics wiring
# --------------------------------------------------------------------------


def test_stats_exclude_missing_values(dataset):
    columns = dataset.wide_columns_map()
    stats = describe(columns["inflation_pct"], variable="inflation_pct")
    assert stats.count == 10
    assert stats.missing == 2
    assert stats.missing_pct == pytest.approx(16.67, abs=0.01)
